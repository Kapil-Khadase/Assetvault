import io
import csv
import datetime
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from services.valuation_engine import get_portfolio_summary
from services.market_data_service import get_latest_market_rates
from services.transaction_service import get_user_transactions

def generate_pdf_report(user_id: int = None, user_info: dict = None, target_currency: str = "INR", portfolio_data: dict = None) -> bytes:
    if portfolio_data is not None:
        portfolio = portfolio_data
    else:
        portfolio = get_portfolio_summary(user_id, target_currency)
    curr_sym = portfolio.get("currency_symbol") or ("₹" if target_currency == "INR" else "$")
    user_info = user_info or {}

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=22,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=15
    )
    section_style = ParagraphStyle(
        'SectionHeader',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        textColor=colors.HexColor('#D97706'),
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'BodyText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        textColor=colors.HexColor('#334155')
    )
    table_hdr_style = ParagraphStyle(
        'TableHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=colors.white
    )
    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        textColor=colors.HexColor('#1E293B')
    )

    story = []

    # 1. Header Banner
    story.append(Paragraph("AssetVault &mdash; Physical Asset Valuation Statement", title_style))
    date_str = datetime.datetime.now().strftime("%B %d, %Y - %H:%M")
    owner_name = user_info.get("name", "Investor")
    story.append(Paragraph(f"Portfolio Owner: <b>{owner_name}</b> | Generated On: <b>{date_str}</b> | Base Currency: <b>{target_currency} ({curr_sym})</b>", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#D97706'), spaceAfter=15))

    # 2. Executive Portfolio KPI Summary
    story.append(Paragraph("1. Executive Portfolio Summary", section_style))

    summary = portfolio.get("summary") if isinstance(portfolio.get("summary"), dict) else portfolio
    total_assets = portfolio.get("total_assets_count", summary.get("total_assets_count", len(portfolio.get("assets", []))))
    active_assets = portfolio.get("total_active_assets", summary.get("total_active_assets", total_assets))
    invested_cost = float(portfolio.get("total_invested_cost", summary.get("total_invested_cost", 0.0)) or 0.0)
    current_val = float(portfolio.get("total_current_market_value", summary.get("total_current_market_value", 0.0)) or 0.0)
    unrealized_pnl = float(portfolio.get("total_unrealized_pnl", summary.get("total_unrealized_pnl", 0.0)) or 0.0)
    overall_roi = float(portfolio.get("overall_roi_percentage", portfolio.get("overall_roi_pct", summary.get("overall_roi_percentage", summary.get("overall_roi_pct", 0.0)))) or 0.0)
    realized_pnl = float(portfolio.get("total_realized_pnl", summary.get("total_realized_pnl", 0.0)) or 0.0)
    sales_count = int(portfolio.get("total_sales_count", summary.get("total_sales_count", 0)) or 0)
    
    kpi_data = [
        [
            Paragraph("<b>Total Assets</b>", table_cell_style),
            Paragraph(f"{active_assets} Active / {total_assets} Total", table_cell_style),
            Paragraph("<b>Total Acquisition Cost</b>", table_cell_style),
            Paragraph(f"{curr_sym}{invested_cost:,.2f}", table_cell_style)
        ],
        [
            Paragraph("<b>Current Estimated Value</b>", table_cell_style),
            Paragraph(f"<b>{curr_sym}{current_val:,.2f}</b>", table_cell_style),
            Paragraph("<b>Unrealized P&L (ROI)</b>", table_cell_style),
            Paragraph(f"{curr_sym}{unrealized_pnl:,.2f} ({overall_roi:+.2f}%)", table_cell_style)
        ],
        [
            Paragraph("<b>Total Realized Profit</b>", table_cell_style),
            Paragraph(f"{curr_sym}{realized_pnl:,.2f} ({sales_count} sales)", table_cell_style),
            Paragraph("<b>Valuation Status</b>", table_cell_style),
            Paragraph("Live Market Linked", table_cell_style)
        ]
    ]

    kpi_table = Table(kpi_data, colWidths=[130, 140, 130, 140])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#E2E8F0')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('PADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 15))

    # 3. Asset Breakdown Table
    story.append(Paragraph("2. Individual Asset Holdings & Market Valuation", section_style))
    
    asset_headers = [
        Paragraph("Asset Name", table_hdr_style),
        Paragraph("Owner / Buyer", table_hdr_style),
        Paragraph("Metal / Type", table_hdr_style),
        Paragraph("Purity", table_hdr_style),
        Paragraph("Net Wt (g)", table_hdr_style),
        Paragraph("Acq Cost", table_hdr_style),
        Paragraph("Est. Value", table_hdr_style),
        Paragraph("P&L", table_hdr_style),
        Paragraph("ROI %", table_hdr_style),
        Paragraph("Status", table_hdr_style)
    ]
    
    table_rows = [asset_headers]
    for a in portfolio.get("assets", []):
        current_val = float(a.get('current_estimated_value') or a.get('current_valuation') or 0.0)
        acq_cost = float(a.get('total_acquisition_cost') or a.get('purchase_price') or 0.0)
        pnl_val = float(a.get('profit_loss') if a.get('profit_loss') is not None else (a.get('unrealized_gain') if a.get('unrealized_gain') is not None else (current_val - acq_cost)))
        pnl_text = f"+{curr_sym}{pnl_val:,.0f}" if pnl_val >= 0 else f"-{curr_sym}{abs(pnl_val):,.0f}"
        roi_val = float(a.get('roi_percentage') if a.get('roi_percentage') is not None else (((pnl_val / acq_cost) * 100) if acq_cost > 0 else 0.0))
        net_wt = float(a.get('net_weight') or a.get('gross_weight') or 0.0)
        
        table_rows.append([
            Paragraph(a.get("asset_name", ""), table_cell_style),
            Paragraph(a.get("owner_name") or a.get("buyer_name") or "Self", table_cell_style),
            Paragraph(f"{a.get('metal', '')} {a.get('asset_type', '')}", table_cell_style),
            Paragraph(str(a.get("purity", "")), table_cell_style),
            Paragraph(f"{net_wt:.2f}", table_cell_style),
            Paragraph(f"{curr_sym}{acq_cost:,.0f}", table_cell_style),
            Paragraph(f"{curr_sym}{current_val:,.0f}", table_cell_style),
            Paragraph(pnl_text, table_cell_style),
            Paragraph(f"{roi_val:+.1f}%", table_cell_style),
            Paragraph(a.get("status", "Active"), table_cell_style),
        ])

    asset_table = Table(table_rows, colWidths=[85, 65, 65, 38, 42, 58, 58, 50, 40, 39])
    asset_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
    ]))
    story.append(asset_table)
    story.append(Spacer(1, 20))

    # 4. Legal Disclaimer / Valuation Notes
    story.append(Paragraph("<b>Valuation Methodology & Disclaimer:</b> Valuations are calculated using latest bullion spot rates adjusted for purity karats and diamond 4Cs attribute models. Figures represent estimated replacement/market benchmarks and do not guarantee resale liquidation prices.", ParagraphStyle('Disc', parent=styles['Italic'], fontSize=7.5, textColor=colors.HexColor('#64748B'))))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()

def generate_excel_report(user_id: int = None, user_info: dict = None, target_currency: str = "INR", portfolio_data: dict = None, transactions_data: list = None) -> bytes:
    if portfolio_data is not None:
        portfolio = portfolio_data
    else:
        portfolio = get_portfolio_summary(user_id, target_currency)
    
    if transactions_data is not None:
        transactions = transactions_data
    elif user_id:
        transactions = get_user_transactions(user_id)
    else:
        transactions = []

    curr_sym = portfolio.get("currency_symbol") or ("₹" if target_currency == "INR" else "$")
    user_info = user_info or {}

    wb = openpyxl.Workbook()

    # Style presets
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    gold_fill = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
    title_font = Font(name="Segoe UI", size=16, bold=True, color="0F172A")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    regular_font = Font(name="Segoe UI", size=10)
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # --- Sheet 1: Portfolio Summary ---
    ws1 = wb.active
    ws1.title = "Portfolio Overview"
    ws1.views.sheetView[0].showGridLines = True

    ws1["A1"] = "AssetVault - Physical Asset Portfolio Report"
    ws1["A1"].font = title_font
    ws1["A2"] = f"Owner: {user_info.get('name', 'User')} | Date: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')} | Currency: {target_currency} ({curr_sym})"
    ws1["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="64748B")

    summary = portfolio.get("summary") if isinstance(portfolio.get("summary"), dict) else portfolio
    kpis = [
        ("Metric", "Value"),
        ("Total Active Holdings", portfolio.get("total_active_assets", summary.get("total_active_assets", 0))),
        ("Total Recorded Assets", portfolio.get("total_assets_count", summary.get("total_assets_count", 0))),
        (f"Total Acquisition Cost ({curr_sym})", portfolio.get("total_invested_cost", summary.get("total_invested_cost", 0.0))),
        (f"Current Estimated Market Value ({curr_sym})", portfolio.get("total_current_market_value", summary.get("total_current_market_value", 0.0))),
        (f"Total Unrealized P&L ({curr_sym})", portfolio.get("total_unrealized_pnl", summary.get("total_unrealized_pnl", 0.0))),
        ("Unrealized ROI (%)", f"{portfolio.get('overall_roi_percentage', portfolio.get('overall_roi_pct', summary.get('overall_roi_percentage', summary.get('overall_roi_pct', 0.0)))):+.2f}%"),
        (f"Total Realized Profit ({curr_sym})", portfolio.get("total_realized_pnl", summary.get("total_realized_pnl", 0.0))),
    ]

    for row_idx, (m, val) in enumerate(kpis, start=4):
        ws1.cell(row=row_idx, column=1, value=m)
        ws1.cell(row=row_idx, column=2, value=val)
        if row_idx == 4:
            ws1.cell(row=row_idx, column=1).fill = gold_fill
            ws1.cell(row=row_idx, column=2).fill = gold_fill
            ws1.cell(row=row_idx, column=1).font = header_font
            ws1.cell(row=row_idx, column=2).font = header_font
        else:
            ws1.cell(row=row_idx, column=1).font = bold_font
            ws1.cell(row=row_idx, column=2).font = regular_font
        ws1.cell(row=row_idx, column=1).border = border_thin
        ws1.cell(row=row_idx, column=2).border = border_thin

    # --- Sheet 2: Asset Inventory ---
    ws2 = wb.create_sheet(title="Asset Inventory")
    ws2.views.sheetView[0].showGridLines = True
    
    headers = [
        "Asset ID", "Asset Name", "Owner / Buyer", "Category", "Type", "Metal", "Purity", 
        "Gross Wt (g)", "Net Wt (g)", "Purchase Date", "Acquisition Cost",
        "Market Rate Used", "Current Estimated Value", "P&L", "ROI %", "Status", "Hallmark / HUID", "Jeweller"
    ]
    
    for col_idx, h in enumerate(headers, start=1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border_thin

    for row_idx, a in enumerate(portfolio.get("assets", []), start=2):
        roi_val = a.get('roi_percentage', 0)
        roi_str = f"{roi_val:+.2f}%" if isinstance(roi_val, (int, float)) else str(roi_val)
        row_data = [
            a.get("id"),
            a.get("asset_name"),
            a.get("owner_name", "-") or "Self",
            a.get("category_name", "General"),
            a.get("asset_type"),
            a.get("metal"),
            a.get("purity"),
            a.get("gross_weight", 0),
            a.get("net_weight", 0),
            a.get("purchase_date"),
            a.get("total_acquisition_cost", 0),
            a.get("market_rate_used", 0),
            a.get("current_estimated_value", 0),
            a.get("profit_loss", 0),
            roi_str,
            a.get("status"),
            a.get("hallmark_huid", "-"),
            a.get("jeweller_name", "-")
        ]
        for col_idx, val in enumerate(row_data, start=1):
            c = ws2.cell(row=row_idx, column=col_idx, value=val)
            c.font = regular_font
            c.border = border_thin

    # --- Sheet 3: Transaction History ---
    ws3 = wb.create_sheet(title="Transactions")
    ws3.views.sheetView[0].showGridLines = True
    tx_headers = ["Tx ID", "Asset Name", "Owner / Holder", "Type", "Date", "Amount", "Realized P&L", "Counterparty", "Notes"]
    for col_idx, h in enumerate(tx_headers, start=1):
        c = ws3.cell(row=1, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.border = border_thin

    for row_idx, t in enumerate(transactions, start=2):
        t_data = [
            t.get("id"),
            t.get("asset_name"),
            t.get("owner_name", "-") or "Self",
            t.get("transaction_type"),
            t.get("transaction_date"),
            t.get("amount"),
            t.get("realized_profit"),
            t.get("counterparty", "-"),
            t.get("notes", "-")
        ]
        for col_idx, val in enumerate(t_data, start=1):
            c = ws3.cell(row=row_idx, column=col_idx, value=val)
            c.font = regular_font
            c.border = border_thin

    # Auto-fit columns
    for ws in [ws1, ws2, ws3]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_csv_export(user_id: int = None, target_currency: str = "INR", portfolio_data: dict = None) -> str:
    if portfolio_data is not None:
        portfolio = portfolio_data
    else:
        portfolio = get_portfolio_summary(user_id, target_currency)
    output = io.StringIO()
    writer = csv.writer(output)
    
    writer.writerow([
        "Asset ID", "Asset Name", "Owner / Holder", "Category", "Type", "Metal", "Purity", 
        "Gross Weight (g)", "Net Weight (g)", "Purchase Date", "Total Acquisition Cost",
        "Current Market Rate", "Current Estimated Value", "Profit Loss", "ROI Percentage", "Status", "HUID"
    ])

    for a in portfolio.get("assets", []):
        current_val = float(a.get('current_estimated_value') or a.get('current_valuation') or 0.0)
        acq_cost = float(a.get('total_acquisition_cost') or a.get('purchase_price') or 0.0)
        pnl_val = float(a.get('profit_loss') if a.get('profit_loss') is not None else (a.get('unrealized_gain') if a.get('unrealized_gain') is not None else (current_val - acq_cost)))
        roi_val = float(a.get('roi_percentage') if a.get('roi_percentage') is not None else (((pnl_val / acq_cost) * 100) if acq_cost > 0 else 0.0))
        net_wt = float(a.get('net_weight') or a.get('gross_weight') or 0.0)
        gross_wt = float(a.get('gross_weight') or a.get('net_weight') or 0.0)

        writer.writerow([
            a.get("id"),
            a.get("asset_name"),
            a.get("owner_name", "") or a.get("buyer_name", "") or "Self",
            a.get("category_name", ""),
            a.get("asset_type", ""),
            a.get("metal", ""),
            a.get("purity", ""),
            gross_wt,
            net_wt,
            a.get("purchase_date", ""),
            acq_cost,
            a.get("market_rate_used", 0),
            current_val,
            pnl_val,
            roi_val,
            a.get("status", ""),
            a.get("hallmark_huid", "")
        ])

    return output.getvalue()
